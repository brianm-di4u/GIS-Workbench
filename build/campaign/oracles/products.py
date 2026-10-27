"""Product oracle (Phase 4): describe every product file the app wrote under a directory with GDAL
(raster chips: size, CRS, corners, RPC grid) and the document readers (PDF pages, PPTX slides,
DOCX paragraphs, KMZ/GPKG layers). Writes <dir>/_products_truth.json.

    micromamba run -n gw-te python -I build/campaign/oracles/products.py <dir>
"""
import json
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common  # noqa: E402,F401
import raster  # noqa: E402
import vector  # noqa: E402


def doc(p):
    s = p.suffix.lower()
    if s == '.pdf':
        from pypdf import PdfReader
        r = PdfReader(str(p))
        pg = r.pages[0]
        return {'kind': 'pdf', 'pages': len(r.pages), 'geo': '/VP' in pg or '/LGIDict' in pg, 'mediabox': [float(v) for v in pg.mediabox]}
    if s == '.pptx':
        from pptx import Presentation
        pr = Presentation(str(p))
        return {'kind': 'pptx', 'slides': len(pr.slides), 'pictures': sum(1 for sl in pr.slides for sh in sl.shapes if sh.shape_type == 13)}
    if s == '.docx':
        from docx import Document
        d = Document(str(p))
        return {'kind': 'docx', 'paragraphs': len(d.paragraphs), 'tables': len(d.tables), 'images': len(d.inline_shapes)}
    return None


def main(d):
    d = Path(d)
    out = {}
    for p in sorted(d.rglob('*')):
        if not p.is_file() or p.name.startswith('_'):
            continue
        try:
            s = p.suffix.lower()
            if s == '.zip':
                with zipfile.ZipFile(p) as z:
                    z.extractall(p.with_suffix(''))
                for q in sorted(p.with_suffix('').iterdir()):
                    if q.suffix.lower() in ('.png', '.jpg'):
                        out[str(q.relative_to(d))] = {'kind': 'raster', **raster.describe(q, stats=False)}
            elif s in ('.tif', '.ntf', '.nitf'):
                out[str(p.relative_to(d))] = {'kind': 'raster', **raster.describe(p, stats=False)}
            elif s in ('.kmz', '.kml', '.gpkg', '.geojson'):
                out[str(p.relative_to(d))] = {'kind': 'vector', **vector.describe(p)}
            else:
                r = doc(p)
                if r:
                    out[str(p.relative_to(d))] = r
        except Exception as e:  # noqa: BLE001 - an unreadable product is the finding
            out[str(p.relative_to(d))] = {'error': f'{type(e).__name__}: {e}'}
    (d / '_products_truth.json').write_text(json.dumps(common._finite(out), indent=1, default=common._jsonable), encoding='utf-8')
    print(json.dumps({k: (v.get('error') or v.get('kind')) for k, v in out.items()}, indent=1))


if __name__ == '__main__':
    main(sys.argv[1])
